from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from docx import Document

from cipher_genius.api.main import app
from cipher_genius.api.knowledge_ingestion_service import KnowledgeIngestionService
from cipher_genius.api.mas_service import MASOrchestrationService
from cipher_genius.api.schemas import AttackDecisionPayload, ExpertGateDecisionPayload
from cipher_genius.core.mas_runtime_support import MASRuntimeSupport
from cipher_genius.features.compliance_reporter import ComplianceStandard
from cipher_genius.models.requirement import (
    PerformanceConstraint,
    PlatformType,
    Requirement,
    ResourceLevel,
    SchemeType,
    SecurityRequirement,
    TargetPlatform,
)
from cipher_genius.models.scheme import (
    CryptographicScheme,
    SchemeArchitecture,
    SchemeMetadata,
    SchemeParameters,
    SecurityAnalysis,
)


client = TestClient(app)

STABLE_ACTORS = {
    "Supervisor",
    "Requirement Analyst",
    "Cryptography Architect",
    "Architect",
    "Security & Compliance Auditor",
    "Code Engineer",
}

EXPECTED_TRACE = [
    "analyst",
    "context_builder",
    "architect",
    "audit",
    "engineer",
    "target_deployer",
    "attack_executor",
    "vulnerability_evaluation",
    "patch_reflection",
    "delivery",
]


def contains_chinese(text: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in str(text))


def build_requirement(*, quantum_resistant: bool = False, security_level: int = 256) -> Requirement:
    return Requirement(
        description="test",
        scheme_type=SchemeType.AUTHENTICATED_ENCRYPTION,
        target_platform=TargetPlatform(type=PlatformType.SERVER, resource_level=ResourceLevel.MODERATE),
        security=SecurityRequirement(security_level=security_level, quantum_resistant=quantum_resistant),
        performance=PerformanceConstraint(),
    )


def build_scheme(service: MASOrchestrationService, name: str, component_names: list[str], *, score: float) -> CryptographicScheme:
    components = [service.component_library.get(item) for item in component_names]
    clean = [item for item in components if item is not None]
    return CryptographicScheme(
        metadata=SchemeMetadata(
            name=name,
            scheme_type=SchemeType.AUTHENTICATED_ENCRYPTION,
            generated_at=datetime.now(),
        ),
        requirements=build_requirement(),
        architecture=SchemeArchitecture(
            components=clean,
            composition={"pattern": "test"},
            dataflow=["Key agreement", "AEAD encryption"],
        ),
        parameters=SchemeParameters(key_size=256, nonce_size=96, tag_size=128),
        security_analysis=SecurityAnalysis(
            threat_model={"adversary": "adaptive"},
            properties=["confidentiality", "integrity"],
            assumptions=["Unique nonce policy"],
            concerns=[],
        ),
        design_rationale="test",
        score=score,
    )


def test_mas_execute_endpoint():
    payload = {
        "requirement": "为支付系统设计兼顾性能与后量子安全的加密方案，延迟目标 20ms。",
        "num_variants": 3,
        "generate_code": True,
        "max_audit_rounds": 4,
    }
    response = client.post("/api/v1/mas/execute", json=payload)
    assert response.status_code == 200

    data = response.json()
    assert data["request_id"]
    assert data["analyst"]["structured_spec"]
    assert isinstance(data["architect"]["candidates"], list)
    assert isinstance(data["auditor_rounds"], list)
    assert isinstance(data["discussion_log"], list)
    assert "status" in data["delivery"]
    assert "scoring" in data["delivery"]
    assert "credibility_assessment" in data
    assert "credibility_assessment" in data["delivery"]
    assert "production_guide" in data["delivery"]
    assert "variant_comparison" in data["delivery"]
    assert "credibility_assessment" in (data["final_scheme"] or {})
    assert data["delivery"]["engine"] == "langgraph"
    assert data["delivery"]["engine_mode"] == "graph-native"
    assert data["delivery"]["attack_loop"]["target_service"]["service_id"]
    assert data["delivery"]["attack_loop"]["target_service"]["status"] == "deployed"
    assert data["delivery"]["attack_loop"]["attack_decision"]["action"] == "execute"
    assert data["delivery"]["attack_loop"]["attack_decision"]["selected_attack_family"]
    assert data["delivery"]["attack_loop"]["attack_results"][0]["status"] == "executed"
    assert data["delivery"]["attack_loop"]["vulnerability_verdict"]["summary"]
    assert data["delivery"]["attack_loop"]["vulnerability_verdict"]["severity"]
    assert data["delivery"]["attack_loop"]["vulnerability_verdict"]["remediation_priority"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["strategy"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["summary"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["rationale"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["implementation_notes"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["validation_steps"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["rollback_notes"]
    assert data["delivery"]["attack_loop"]["patch_spec"]["regression_focus"]
    assert (
        data["delivery"]["attack_loop"]["patch_execution"]["patch_id"]
        == data["delivery"]["attack_loop"]["patch_spec"]["patch_id"]
    )
    assert data["delivery"]["attack_loop"]["patch_execution"]["status"] in {"applied", "validated"}
    assert data["delivery"]["attack_loop"]["patch_execution"]["workspace"]
    assert data["delivery"]["attack_loop"]["patch_execution"]["execution_contract_version"] == "v1"
    assert data["delivery"]["attack_loop"]["patch_execution"]["validation_results"]
    assert data["delivery"]["attack_loop"]["patch_execution"]["validation_summary"]["passed"] >= 1
    assert data["delivery"]["attack_loop"]["patch_execution"]["validation_results"][0]["step_kind"]
    assert data["delivery"]["attack_loop"]["patch_execution"]["changed_artifact_summaries"]
    assert data["delivery"]["attack_loop"]["patch_execution"]["artifact_inventory"]["changed_code_artifact_count"] >= 1
    assert data["delivery"]["attack_loop"]["patch_execution"]["patch_artifact_refs"]
    assert data["delivery"]["attack_loop"]["patch_execution"]["rollback_artifact_refs"]
    assert len(data["delivery"]["attack_loop"]["rounds"]) >= 2
    assert data["delivery"]["attack_loop"]["rounds"][0]["attack_decision"]["action"] == "execute"
    assert data["delivery"]["attack_loop"]["rounds"][0]["vulnerability_verdict"]["summary"]
    assert data["delivery"]["attack_loop"]["rounds"][1]["round_kind"] == "regression"
    assert data["delivery"]["attack_loop"]["regression_attack_decision"]["action"] == "replan"
    assert data["delivery"]["attack_loop"]["regression_vulnerability_verdict"]["summary"]
    assert data["delivery"]["attack_loop"]["regression_vulnerability_verdict"]["severity"]
    assert data["delivery"]["attack_loop"]["reflection_cards"]
    assert data["delivery"]["attack_loop"]["reflection_cards"][0]["card_type"] == "reflection"
    assert data["delivery"]["attack_loop"]["reflection_cards"][0]["prompt_changes"]
    assert data["delivery"]["attack_loop"]["reflection_cards"][1]["card_type"] == "regression_summary"
    assert data["delivery"]["attack_loop"]["rounds"][1]["attack_decision"]["action"] == "replan"
    assert data["delivery"]["attack_loop"]["rounds"][1]["vulnerability_verdict"]["summary"]
    assert data["delivery"]["attack_loop"]["rounds"][1]["mode"] == "executed"
    assert data["delivery"]["attack_loop"]["rounds"][1]["patch_applied"] is True
    assert data["delivery"]["attack_loop"]["rounds"][1]["patch_execution"]["status"] in {"applied", "validated"}
    assert data["delivery"]["attack_loop"]["rounds"][1]["attack_results"]
    assert data["delivery"]["attack_loop"]["regression_attack_results"][0]["status"] == "executed"
    assert data["delivery"]["context_projections"]["generation"]["agent_id"] == "generation_agent"
    assert data["delivery"]["context_projections"]["audit"]["agent_id"] == "audit_agent"
    assert data["delivery"]["context_projections"]["attack_planning"]["agent_id"] == "attack_planning_agent"
    assert (
        data["delivery"]["context_projections"]["vulnerability_evaluation"]["agent_id"]
        == "vulnerability_agent"
    )
    assert data["delivery"]["context_projections"]["expert_gate"]["agent_id"] == "expert_gate_agent"
    assert data["delivery"]["context_projections"]["patch"]["agent_id"] == "patch_agent"
    assert data["delivery"]["context_projections"]["reflection"]["agent_id"] == "reflection_agent"
    assert data["delivery"]["context_projections"]["generation"]["constraints"]
    assert any(
        item["card_type"] == "generation_runtime_input"
        and item["payload"].get("parsed_requirement", {}).get("requirement")
        and item.get("card_family") == "runtime_input"
        and item.get("card_contract_version") == "v1"
        and item.get("lineage_ref")
        and item.get("payload", {}).get("typed_contract", {}).get("contract_ref")
        == "runtime_input:generation_runtime_input:v1"
        and item.get("payload", {}).get("ref_lookup_hint", {}).get("lookup_strategy")
        == "card_refs_then_projection_then_handoff_then_replay"
        for item in data["delivery"]["context_projections"]["generation"]["cards"]
    )
    assert any(
        item["card_type"] == "audit_runtime_input"
        and item["payload"].get("schemes")
        for item in data["delivery"]["context_projections"]["audit"]["cards"]
    )
    assert any(
        item["card_type"] == "audit_runtime_input"
        and item["payload"].get("scheme_entries")
        for item in data["delivery"]["context_projections"]["audit"]["cards"]
    )
    assert any(
        item["card_type"] == "audit_decision_input"
        and item["payload"].get("proposal_id")
        and item["payload"].get("compliance_score") is not None
        and item["payload"].get("risk_score") is not None
        for item in data["delivery"]["context_projections"]["audit"]["cards"]
    )
    assert data["delivery"]["context_projections"]["attack_planning"]["artifact_refs"]
    assert any(
        item["card_type"] == "attack_decision"
        and item.get("card_family") == "decision"
        for item in data["delivery"]["context_projections"]["vulnerability_evaluation"]["cards"]
    )
    assert any(
        item["card_type"] == "attack_result_summary"
        for item in data["delivery"]["context_projections"]["vulnerability_evaluation"]["cards"]
    )
    assert any(
        item["card_type"] == "attack_result_summary"
        and item["payload"].get("attack_result_summaries")
        for item in data["delivery"]["context_projections"]["vulnerability_evaluation"]["cards"]
    )
    assert any(
        item["card_type"] == "vulnerability_verdict"
        for item in data["delivery"]["context_projections"]["patch"]["cards"]
    )
    assert any(
        item["card_type"] == "expert_gate_decision"
        and item["payload"].get("action")
        and item["payload"].get("decision_family")
        and item["payload"].get("route_target")
        for item in data["delivery"]["context_projections"]["patch"]["cards"]
    )
    assert any(
        item["card_type"] == "attack_artifact_summary"
        and item["payload"].get("attack_result_summaries")
        for item in data["delivery"]["context_projections"]["patch"]["cards"]
    )
    assert data["delivery"]["attack_loop"]["expert_gate_decision"]["action"] in {
        "patch_required",
        "patch_required_with_regression",
        "retry_attack",
        "observe_only",
    }
    assert data["delivery"]["attack_loop"]["expert_gate_decision"]["decision_family"] in {
        "patch_flow",
        "retry_flow",
        "observation_flow",
    }
    assert data["delivery"]["attack_loop"]["expert_gate_decision"]["route_target"] in {
        "patch_agent",
        "attack_planning_agent",
        "delivery",
    }
    assert data["delivery"]["attack_loop"]["expert_gate_decision"]["residual_risk_summary"]
    assert any(
        item["card_type"] == "regression_verdict_input"
        for item in data["delivery"]["context_projections"]["reflection"]["cards"]
    )
    assert any(
        item["card_type"] == "patch_execution"
        and item["payload"].get("validation_results")
        and item["payload"].get("validation_summary")
        and item.get("card_family") == "patch"
        for item in data["delivery"]["context_projections"]["reflection"]["cards"]
    )
    reflection_memory_cards = [
        item
        for item in data["delivery"]["context_projections"]["generation"]["cards"]
        if item.get("card_type") == "reflection_memory"
    ]
    if reflection_memory_cards:
        assert any(
            item.get("card_family") == "reflection"
            and item.get("card_contract_version") == "v1"
            and item.get("lineage_ref")
            for item in reflection_memory_cards
        )
    assert any(
        item["card_type"] == "patch_artifact_summary"
        and item["payload"].get("changed_artifact_summaries")
        and item["payload"]["changed_artifact_summaries"][0].get("diff_preview")
        for item in data["delivery"]["context_projections"]["reflection"]["cards"]
    )
    assert len(data["delivery"]["memory_handoffs"]) >= 6
    assert data["delivery"]["memory_handoffs"][0]["projection"]["agent_id"] == "generation_agent"
    assert data["delivery"]["memory_handoffs"][-1]["projection"]["agent_id"] == "reflection_agent"
    assert data["delivery"]["sandbox_dispatcher"]["backend"] == "local-dispatcher"
    assert data["delivery"]["sandbox_dispatcher"]["policy"]["policy_id"] == "local-sandbox-default"
    assert data["delivery"]["sandbox_dispatcher"]["baseline_deployment"]["decision"] == "approved"
    assert data["delivery"]["sandbox_dispatcher"]["baseline_attack"]["status"] == "executed"
    assert data["delivery"]["sandbox_dispatcher"]["patch_apply"]["status"] == "executed"
    assert data["delivery"]["sandbox_dispatcher"]["regression_deployment"]["decision"] == "approved"
    assert data["delivery"]["sandbox_dispatcher"]["regression_attack"]["status"] == "executed"
    assert data["delivery"]["sandbox_dispatcher"]["rollback_plan"]["status"] == "executed"
    assert data["delivery"]["sandbox_dispatcher"]["baseline_attack"]["failure_items"] == []
    assert (
        data["delivery"]["attack_loop"]["patch_execution"]["patch_dispatch_id"]
        == data["delivery"]["sandbox_dispatcher"]["patch_apply"]["dispatch_id"]
    )
    assert (
        data["delivery"]["attack_loop"]["patch_execution"]["rollback_dispatch_id"]
        == data["delivery"]["sandbox_dispatcher"]["rollback_plan"]["dispatch_id"]
    )
    assert data["delivery"]["sandbox_dispatcher"]["baseline_attack"]["audit_trail"][-1]["event_kind"] == "executed"
    assert data["delivery"]["sandbox_dispatcher"]["regression_attack"]["audit_trail"][-1]["event_kind"] == "executed"
    assert data["delivery"]["backend_architecture"]["control_plane"]["summary"]["stage_count"] == len(EXPECTED_TRACE)
    assert data["delivery"]["backend_architecture"]["control_plane"]["summary"]["contract_version"] == "v1"
    assert data["delivery"]["backend_architecture"]["control_plane"]["summary"]["retryable_stage_count"] >= 1
    attack_control = next(
        item
        for item in data["delivery"]["backend_architecture"]["control_plane"]["invocations"]
        if item["stage"] == "attack_executor"
    )
    assert attack_control["checkpoint_ref"]
    assert attack_control["result_ref"]
    assert attack_control["retryable"] is True
    assert attack_control["retry_strategy"] == "replan_or_handoff_before_skip"
    assert attack_control["termination_mode"] == "execute"
    assert attack_control["decision_source"] == "attack_decision.action"
    assert attack_control["contract"]["failure_action"] == "skip_dispatch_or_handoff"
    assert attack_control["contract"]["next_stages"] == ["vulnerability_evaluation"]
    assert attack_control["contract"]["retryable"] is True
    assert attack_control["contract"]["retry_strategy"] == "replan_or_handoff_before_skip"
    assert attack_control["contract"]["decision_source"] == "attack_decision.action"
    assert attack_control["contract"]["termination_conditions"]
    attack_result_envelope = next(
        item
        for item in data["delivery"]["backend_architecture"]["control_plane"]["results"]
        if item["stage"] == "attack_executor"
    )
    assert attack_result_envelope["decision_signal"] == "execute"
    assert attack_result_envelope["decision_source"] == "attack_decision.action"
    assert attack_result_envelope["termination_signal"] == "execute"
    assert attack_result_envelope["failure_action"] == "skip_dispatch_or_handoff"
    assert data["delivery"]["backend_architecture"]["memory_bus"]["projection_count"] >= 6
    assert data["delivery"]["backend_architecture"]["memory_bus"]["summary"]["typed_family_count"] >= 4
    assert data["delivery"]["backend_architecture"]["memory_bus"]["summary"]["typed_contract_count"] >= 4
    assert data["delivery"]["backend_architecture"]["memory_bus"]["summary"]["family_counts"]["reflection"] >= 1
    assert data["delivery"]["backend_architecture"]["memory_bus"]["typed_families"]
    assert data["delivery"]["backend_architecture"]["memory_bus"]["typed_contracts"]
    assert data["delivery"]["backend_architecture"]["memory_bus"]["replay_snapshot_card"]["typed_contract_counts"]
    assert data["delivery"]["backend_architecture"]["execution_plane"]["summary"]["stage_count"] >= 4
    assert data["delivery"]["backend_architecture"]["execution_plane"]["summary"]["contract_version"] == "v1"
    assert data["delivery"]["backend_architecture"]["execution_plane"]["summary"]["operation_count"] >= 6
    assert data["delivery"]["backend_architecture"]["execution_plane"]["plan"]["executor_kinds"]
    assert data["delivery"]["backend_architecture"]["execution_plane"]["plan"]["executor_matrix"]
    assert data["delivery"]["backend_architecture"]["execution_plane"]["summary"]["executor_matrix_count"] >= 3
    assert (
        data["delivery"]["backend_architecture"]["execution_plane"]["summary"]["operation_kind_counts"]["patch_apply"]
        >= 1
    )
    assert data["delivery"]["backend_architecture"]["execution_plane"]["plan"]["operation_kinds"]
    assert any(
        item["operation_kind"] == "regression_replay" and item["depends_on"] and item["output_refs"]
        for item in data["delivery"]["backend_architecture"]["execution_plane"]["operations"]
    )
    assert any(
        item["operation_kind"] == "rollback" and item["status"] == "executed"
        for item in data["delivery"]["backend_architecture"]["execution_plane"]["operations"]
    )
    assert data["delivery"]["backend_architecture"]["replay_plane"]["timeline_summary"]["snapshot_count"] >= 1
    assert data["delivery"]["backend_architecture"]["replay_plane"]["timeline_summary"]["latest_typed_family_count"] >= 1
    assert data["delivery"]["backend_architecture"]["replay_plane"]["timeline_summary"]["latest_typed_contract_count"] >= 1
    assert data["delivery"]["backend_architecture"]["replay_plane"]["timeline_summary"]["latest_dispatch_ref_count"] >= 1
    assert data["delivery"]["backend_architecture"]["replay_plane"]["timeline_summary"]["latest_failed_dispatch_ref_count"] >= 0
    assert isinstance(
        data["delivery"]["backend_architecture"]["replay_plane"]["timeline_summary"][
            "latest_executor_handoff_trace_count"
        ],
        int,
    )
    assert data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["projection_refs"]
    assert data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["dispatch_refs"]
    assert isinstance(
        data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["failed_dispatch_refs"],
        list,
    )
    assert data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["typed_family_counts"]
    assert data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["typed_contract_counts"]
    assert isinstance(
        data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"][
            "executor_handoff_trace_summaries"
        ],
        list,
    )
    assert data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["metadata"]["dispatch_catalog"]
    assert data["delivery"]["backend_architecture"]["replay_plane"]["latest_snapshot"]["metadata"]["patch_execution_artifact_refs"]
    assert data["delivery"]["code_artifacts"]["pseudocode_ready"] is True


def test_mas_execute_endpoint_with_langgraph_native_engine():
    case_id = f"case-api-{uuid4().hex[:8]}"
    payload = {
        "case_id": case_id,
        "requirement": "为建筑项目 CDE 中的 BIM/IFC 交付设计后量子可迁移的低延迟签名与加密方案。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    response = client.post("/api/v1/mas/execute?use_langgraph=true", json=payload)
    assert response.status_code == 200

    data = response.json()
    assert data["delivery"]["engine"] == "langgraph"
    assert data["delivery"]["engine_mode"] == "graph-native"
    assert data["delivery"]["workflow_trace"] == EXPECTED_TRACE
    assert data["analyst"]["structured_spec"]
    assert data["case_id"] == case_id
    assert data["case_memory"]["case_id"] == case_id
    assert data["case_memory"]["recent_reflections"]
    assert data["delivery"]["case_memory_summary"]["case_id"] == case_id
    assert data["delivery"]["case_memory_summary"]["reflection_count"] >= 1
    assert isinstance(data["auditor_rounds"], list)
    assert "evidence_pack" in data
    assert data["evidence_pack"]["backend"] in {"local", "qdrant"}

    second_response = client.post("/api/v1/mas/execute?use_langgraph=true", json=payload)
    assert second_response.status_code == 200
    second = second_response.json()
    assert any(
        item.get("card_type") == "reflection_memory"
        and item.get("payload", {}).get("latest_prompt_changes")
        and item.get("card_family") == "reflection"
        and item.get("card_contract_version") == "v1"
        and item.get("lineage_ref")
        for item in second["delivery"]["context_projections"]["generation"]["cards"]
    )
    assert any(
        item.get("card_type") == "generation_runtime_input"
        and item.get("payload", {}).get("structured_spec")
        for item in second["delivery"]["context_projections"]["generation"]["cards"]
    )
    assert any(
        item.get("card_type") == "audit_runtime_input"
        and item.get("payload", {}).get("schemes")
        for item in second["delivery"]["context_projections"]["audit"]["cards"]
    )
    assert any(
        item.get("card_type") == "audit_runtime_input"
        and item.get("payload", {}).get("scheme_entries")
        for item in second["delivery"]["context_projections"]["audit"]["cards"]
    )
    assert any(
        item.get("card_type") == "audit_decision_input"
        and item.get("payload", {}).get("proposal_id")
        and item.get("payload", {}).get("compliance_score") is not None
        and item.get("payload", {}).get("risk_score") is not None
        for item in second["delivery"]["context_projections"]["audit"]["cards"]
    )
    assert any(
        "[反思回灌]" in str(candidate.get("design_rationale") or "")
        for candidate in second["architect"]["candidates"]
    )
    assert any(
        any("项目级反思" in str(change or "") for change in round_item.get("recommended_changes", []))
        for round_item in second["auditor_rounds"]
    )


def test_case_timeline_endpoints_return_replay_overview():
    case_id = f"case-timeline-{uuid4().hex[:8]}"
    payload = {
        "case_id": case_id,
        "requirement": "Design a post-quantum trust plan for BIM/IFC delivery through a construction CDE with audit trail and project memory support.",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    execute_response = client.post("/api/v1/mas/execute", json=payload)
    assert execute_response.status_code == 200

    timeline_response = client.get(f"/api/v1/cases/{case_id}/timeline")
    assert timeline_response.status_code == 200
    timeline = timeline_response.json()
    assert timeline["case_id"] == case_id
    assert timeline["timeline_summary"]["snapshot_count"] >= 1
    assert timeline["timeline_summary"]["event_count"] >= 1
    assert timeline["timeline_summary"]["latest_typed_contract_count"] >= 1
    assert timeline["timeline_summary"]["latest_dispatch_ref_count"] >= 1
    assert timeline["timeline_summary"]["latest_failed_dispatch_ref_count"] >= 0
    assert timeline["timeline_summary"]["latest_dispatch_summary_count"] >= 1
    assert timeline["timeline_summary"]["latest_failed_dispatch_summary_count"] >= 0
    assert isinstance(timeline["timeline_summary"]["latest_executor_handoff_trace_count"], int)
    assert timeline["latest_snapshot"]["projection_refs"]
    assert timeline["latest_snapshot"]["dispatch_refs"]
    assert timeline["latest_snapshot"]["dispatch_summaries"]
    assert isinstance(timeline["latest_snapshot"]["failed_dispatch_refs"], list)
    assert isinstance(timeline["latest_snapshot"]["failed_dispatch_summaries"], list)
    assert isinstance(timeline["latest_snapshot"]["executor_handoff_trace_summaries"], list)
    artifact_lookup_refs = timeline["latest_snapshot"]["artifact_lookup_refs"]
    assert artifact_lookup_refs
    scoped_artifact_ref = next(
        (item for item in artifact_lookup_refs if str(item).startswith("artifact:")),
        artifact_lookup_refs[0],
    )

    events_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/events",
        params={"event_kind": "typed_card_contract", "limit": 10},
    )
    assert events_response.status_code == 200
    events = events_response.json()
    assert events["case_id"] == case_id
    assert events["scope"]["event_kind"] == "typed_card_contract"
    assert events["total"] >= 1
    assert all(item["event_kind"] == "typed_card_contract" for item in events["items"])

    projection_events_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/events",
        params={"projection_ref": "generation_agent:generation-r1", "limit": 10},
    )
    assert projection_events_response.status_code == 200
    projection_events = projection_events_response.json()
    assert projection_events["case_id"] == case_id
    assert projection_events["scope"]["projection_ref"] == "generation_agent:generation-r1"
    assert projection_events["total"] >= 1

    artifact_events_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/events",
        params={"artifact_lookup_ref": scoped_artifact_ref, "limit": 10},
    )
    assert artifact_events_response.status_code == 200
    artifact_events = artifact_events_response.json()
    assert artifact_events["case_id"] == case_id
    assert artifact_events["scope"]["artifact_lookup_ref"] == scoped_artifact_ref
    assert artifact_events["total"] >= 1
    assert all(item["run_id"] == timeline["latest_run_id"] for item in artifact_events["items"])

    snapshots_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/snapshots",
        params={"contract_ref": "decision:attack_decision:v1", "limit": 10},
    )
    assert snapshots_response.status_code == 200
    snapshots = snapshots_response.json()
    assert snapshots["case_id"] == case_id
    assert snapshots["scope"]["contract_ref"] == "decision:attack_decision:v1"
    assert snapshots["total"] >= 1
    assert all(
        "decision:attack_decision:v1" in item["typed_contract_counts"]
        for item in snapshots["items"]
    )

    projection_snapshots_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/snapshots",
        params={"projection_ref": "generation_agent:generation-r1", "limit": 10},
    )
    assert projection_snapshots_response.status_code == 200
    projection_snapshots = projection_snapshots_response.json()
    assert projection_snapshots["case_id"] == case_id
    assert projection_snapshots["scope"]["projection_ref"] == "generation_agent:generation-r1"
    assert projection_snapshots["total"] >= 1
    assert all(
        "generation_agent:generation-r1" in item["projection_refs"]
        for item in projection_snapshots["items"]
    )

    artifact_snapshots_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/snapshots",
        params={"artifact_lookup_ref": scoped_artifact_ref, "limit": 10},
    )
    assert artifact_snapshots_response.status_code == 200
    artifact_snapshots = artifact_snapshots_response.json()
    assert artifact_snapshots["case_id"] == case_id
    assert artifact_snapshots["scope"]["artifact_lookup_ref"] == scoped_artifact_ref
    assert artifact_snapshots["total"] >= 1
    assert all(
        scoped_artifact_ref in item["artifact_lookup_refs"]
        for item in artifact_snapshots["items"]
    )

    drilldown_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/drilldown",
        params={"projection_ref": "generation_agent:generation-r1", "limit": 10},
    )
    assert drilldown_response.status_code == 200
    drilldown = drilldown_response.json()
    assert drilldown["case_id"] == case_id
    assert drilldown["scope"]["projection_ref"] == "generation_agent:generation-r1"
    assert drilldown["summary"]["matched_run_ids"]
    assert drilldown["summary"]["snapshot_count"] >= 1
    assert drilldown["summary"]["event_count"] >= 1
    assert "generation_agent:generation-r1" in drilldown["summary"]["projection_refs"]
    assert drilldown["projection_relationships"]
    assert drilldown["projection_relationships"][0]["relation_kind"] == "projection"
    assert drilldown["projection_relationships"][0]["relation_ref"] == "generation_agent:generation-r1"
    assert "generation_agent" in drilldown["projection_relationships"][0]["agent_ids"]
    assert drilldown["projection_relationships"][0]["snapshot_count"] >= 1
    assert drilldown["handoff_relationships"]
    assert drilldown["handoff_relationships"][0]["relation_kind"] == "handoff"
    assert drilldown["handoff_relationships"][0]["target_agent_id"] or drilldown["handoff_relationships"][0]["linked_projection_refs"]
    assert drilldown["service_trajectories"]
    assert drilldown["service_trajectories"][0]["target_service_ref"]
    assert drilldown["service_trajectories"][0]["event_kind_counts"]
    assert drilldown["snapshots"]
    assert drilldown["events"]
    assert drilldown["version_lineage"]

    artifact_drilldown_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/drilldown",
        params={"artifact_lookup_ref": scoped_artifact_ref, "limit": 10},
    )
    assert artifact_drilldown_response.status_code == 200
    artifact_drilldown = artifact_drilldown_response.json()
    assert artifact_drilldown["case_id"] == case_id
    assert artifact_drilldown["scope"]["artifact_lookup_ref"] == scoped_artifact_ref
    assert artifact_drilldown["summary"]["matched_run_ids"]
    assert artifact_drilldown["summary"]["dispatch_refs"]
    assert isinstance(artifact_drilldown["summary"]["failed_dispatch_refs"], list)
    assert artifact_drilldown["summary"]["dispatch_summaries"]
    assert isinstance(artifact_drilldown["summary"]["failed_dispatch_summaries"], list)
    assert isinstance(artifact_drilldown["summary"]["executor_handoff_trace_summaries"], list)
    assert scoped_artifact_ref in artifact_drilldown["summary"]["artifact_lookup_refs"]
    assert artifact_drilldown["events"]
    assert all(item["run_id"] == timeline["latest_run_id"] for item in artifact_drilldown["events"])

    lineage_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/lineage",
        params={"run_id": timeline["latest_run_id"], "baseline_version": "v1", "limit": 10},
    )
    assert lineage_response.status_code == 200
    lineage = lineage_response.json()
    assert lineage["case_id"] == case_id
    assert lineage["total"] >= 1
    assert all(item["baseline_version"] == "v1" for item in lineage["items"])
    assert all(item["run_id"] == timeline["latest_run_id"] for item in lineage["items"])
    assert all(item["target_service_ref"] for item in lineage["items"])
    assert all(item["patch_id"] for item in lineage["items"])


def test_mas_execute_endpoint_deprecated_flag_still_routes_to_langgraph():
    payload = {
        "requirement": "为建筑项目 BIM/IFC 可信交付设计支持审计留痕与后量子迁移的密码方案。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    response = client.post("/api/v1/mas/execute?use_langgraph=false", json=payload)
    assert response.status_code == 200

    data = response.json()
    assert data["delivery"]["engine"] == "langgraph"
    assert data["delivery"]["engine_mode"] == "graph-native"


def test_mas_execute_endpoint_supports_attack_handoff(monkeypatch):
    def fake_plan_attack(self, projection, *, run_id):
        return (
            AttackDecisionPayload(
                decision_id=f"attack-decision-{run_id[:8]}",
                target_service_ref="svc-test",
                action="handoff_to_vulnerability",
                action_label="转交漏洞评估",
                rationale="当前证据已经足以进入漏洞评估。",
                selected_attack_family="oracle_probe",
                expected_outcome="直接形成漏洞裁决。",
                confidence=0.82,
                stop_conditions=["不再新增攻击任务"],
                next_step="handoff_to_vulnerability",
            ),
            [],
        )

    monkeypatch.setattr(MASRuntimeSupport, "_plan_attack_from_projection", fake_plan_attack)

    payload = {
        "requirement": "为支付系统设计兼顾性能与后量子安全的加密方案，延迟目标 20ms。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    response = client.post("/api/v1/mas/execute", json=payload)
    assert response.status_code == 200

    data = response.json()
    attack_loop = data["delivery"]["attack_loop"]
    assert attack_loop["attack_decision"]["action"] == "handoff_to_vulnerability"
    assert attack_loop["attack_results"] == []
    assert attack_loop["rounds"][0]["mode"] == "handoff"
    assert attack_loop["vulnerability_verdict"]["summary"]
    route_target = attack_loop["expert_gate_decision"]["route_target"]
    assert route_target in {"patch_agent", "attack_planning_agent", "delivery"}
    if route_target == "patch_agent":
        assert attack_loop["regression_attack_decision"]["action"] == "handoff_to_vulnerability"
        assert attack_loop["rounds"][1]["mode"] == "handoff"
        assert attack_loop["loop_status"] == "baseline_handoff_regression_handoff"
        assert attack_loop["regression_vulnerability_verdict"]["summary"]
    else:
        assert attack_loop["patch_spec"] == {}
        assert attack_loop["regression_attack_decision"] == {}
        assert len(attack_loop["rounds"]) == 1
        assert attack_loop["loop_status"] in {
            "baseline_handoff_delivery_observation",
            "baseline_handoff_retry_follow_up",
        }
    assert data["delivery"]["sandbox_dispatcher"]["baseline_attack"]["status"] == "skipped"
    assert data["delivery"]["sandbox_dispatcher"]["baseline_attack"]["decision"] == "handoff_to_vulnerability"
    assert data["delivery"]["sandbox_dispatcher"]["baseline_attack"]["audit_trail"][-1]["event_kind"] == "handoff"
    if route_target == "patch_agent":
        assert data["delivery"]["sandbox_dispatcher"]["regression_attack"]["status"] == "skipped"
    else:
        assert data["delivery"]["sandbox_dispatcher"]["regression_attack"] == {}


def test_mas_execute_endpoint_supports_same_run_retry_route(monkeypatch):
    original_plan_attack = MASRuntimeSupport._plan_attack_from_projection
    gate_calls = {"count": 0}
    case_id = f"case-retry-{uuid4().hex[:8]}"

    def fake_expert_gate(self, projection, *, run_id):  # noqa: ANN001
        gate_calls["count"] += 1
        target_service_ref = ""
        for artifact in projection.artifact_refs:
            metadata = artifact.metadata or {}
            target_service_ref = str(metadata.get("service_id") or "").strip()
            if target_service_ref:
                break
        if gate_calls["count"] == 1:
            return ExpertGateDecisionPayload(
            decision_id=f"expert-gate-{run_id[:8]}",
            target_service_ref=target_service_ref or "svc-retry",
            decision_family="retry_flow",
            decision_family_label="补充验证",
            action="retry_attack",
            action_label="补充攻击验证",
            route_target="attack_planning_agent",
            route_target_label="攻击重规划窗口",
            rationale="当前基线证据不足，需要同轮补充攻击验证。",
            confidence=0.8,
            residual_risk_summary="当前残余风险仍需要额外采样确认。",
            follow_up_actions=["扩大攻击采样", "重新确认错误路径"],
        )
        return ExpertGateDecisionPayload(
            decision_id=f"expert-gate-{run_id[:8]}",
            target_service_ref=target_service_ref or "svc-retry",
            decision_family="observation_flow",
            decision_family_label="观察收敛",
            action="observe_only",
            action_label="观察收敛",
            route_target="delivery",
            route_target_label="交付观察窗口",
            rationale="补充攻击证据已经足够，本轮先以观察收口。",
            confidence=0.83,
            residual_risk_summary="残余风险可先在下一轮按反思卡继续跟进。",
            follow_up_actions=["沉淀补充攻击结论", "下一轮继续补强攻击策略"],
        )

    def fake_plan_attack(self, projection, *, run_id):  # noqa: ANN001
        if str(run_id).endswith("-retry"):
            _, baseline_specs = original_plan_attack(self, projection, run_id=run_id)
            return (
                AttackDecisionPayload(
                    decision_id=f"attack-decision-{run_id[:8]}",
                    target_service_ref="svc-retry",
                    action="execute",
                    action_label="执行补充攻击",
                    rationale="专家闸门要求同轮补充验证。",
                    selected_attack_family="oracle_probe",
                    expected_outcome="形成补充漏洞证据。",
                    confidence=0.79,
                    stop_conditions=["补充证据达到阈值"],
                    next_step="dispatch_attack",
                ),
                [
                    item.model_copy(update={"attack_id": f"{item.attack_id}-retry"})
                    for item in baseline_specs
                ],
            )
        return original_plan_attack(self, projection, run_id=run_id)

    monkeypatch.setattr(MASRuntimeSupport, "_decide_expert_gate_from_projection", fake_expert_gate)
    monkeypatch.setattr(MASRuntimeSupport, "_plan_attack_from_projection", fake_plan_attack)

    payload = {
        "case_id": case_id,
        "requirement": "为支付系统设计兼顾性能与后量子安全的加密方案，延迟目标 20ms。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    response = client.post("/api/v1/mas/execute", json=payload)
    assert response.status_code == 200

    data = response.json()
    attack_loop = data["delivery"]["attack_loop"]
    assert attack_loop["expert_gate_decision"]["route_target"] == "attack_planning_agent"
    assert attack_loop["retry_expert_gate_decision"]["route_target"] == "delivery"
    assert attack_loop["patch_spec"] == {}
    assert attack_loop["retry_attack_decision"]["action"] == "execute"
    assert attack_loop["retry_attack_results"]
    assert attack_loop["retry_vulnerability_verdict"]["summary"]
    assert attack_loop["same_run_retry_budget"] == 1
    assert attack_loop["same_run_retry_used"] == 1
    assert attack_loop["same_run_retry_remaining"] == 0
    assert attack_loop["same_run_retry_summary"]["requested"] is True
    assert attack_loop["same_run_retry_summary"]["started"] is True
    assert attack_loop["same_run_retry_summary"]["final_resolution"] == "observed_after_retry"
    assert attack_loop["same_run_retry_summary"]["attempt_count"] == 1
    assert attack_loop["same_run_retry_summary"]["attempt_trace"][0]["status"] == "executed"
    assert attack_loop["same_run_retry_summary"]["attempt_trace"][0]["termination_reason"] == "observed_after_retry"
    assert (
        attack_loop["same_run_retry_summary"]["attempt_trace"][0]["planning_projection_ref"]
        == "attack_planning_agent:attack-plan-r2"
    )
    assert (
        attack_loop["same_run_retry_summary"]["attempt_trace"][0]["vulnerability_projection_ref"]
        == "vulnerability_agent:vulnerability-r2"
    )
    assert (
        attack_loop["same_run_retry_summary"]["attempt_trace"][0]["expert_gate_projection_ref"]
        == "expert_gate_agent:expert-gate-r2"
    )
    assert data["delivery"]["context_projections"]["attack_planning_retry"]["agent_id"] == "attack_planning_agent"
    assert (
        data["delivery"]["context_projections"]["vulnerability_evaluation_retry"]["agent_id"]
        == "vulnerability_agent"
    )
    assert any(
        item["card_type"] == "retry_context_summary"
        for item in data["delivery"]["context_projections"]["attack_planning_retry"]["cards"]
    )
    retry_summary_card = next(
        item
        for item in data["delivery"]["context_projections"]["attack_planning_retry"]["cards"]
        if item["card_type"] == "retry_context_summary"
    )
    assert retry_summary_card["payload"]["compression_policy"] == "retain_summary_and_refs_drop_raw_details"
    assert retry_summary_card["payload"]["retained_refs"]
    assert retry_summary_card["payload"]["dropped_detail_hints"]
    assert retry_summary_card["payload"]["resume_checkpoint_ref"]
    assert retry_summary_card["payload"]["resume_inputs"]["retained_refs"]
    assert any(
        item["card_type"] == "retry_context_summary"
        for item in data["delivery"]["context_projections"]["vulnerability_evaluation_retry"]["cards"]
    )
    memory_bus = data["delivery"]["backend_architecture"]["memory_bus"]
    assert memory_bus["summary"]["retry_window_count"] == 3
    assert memory_bus["summary"]["retry_handoff_count"] >= 3
    assert "attack_planning_agent:attack-plan-r2" in memory_bus["summary"]["retry_projection_refs"]
    assert "vulnerability_agent:vulnerability-r2" in memory_bus["summary"]["retry_projection_refs"]
    assert "expert_gate_agent:expert-gate-r2" in memory_bus["summary"]["retry_projection_refs"]
    assert memory_bus["summary"]["retry_lineage_refs"]
    assert memory_bus["summary"]["retry_typed_contract_refs"]
    assert memory_bus["summary"]["retry_compression_stages"]
    assert memory_bus["summary"]["retry_compression_policies"] == [
        "retain_summary_and_refs_drop_raw_details"
    ]
    assert memory_bus["summary"]["retry_retained_refs"]
    assert memory_bus["summary"]["retry_resume_checkpoint_refs"]
    assert memory_bus["summary"]["retry_resume_input_refs"]
    assert len(attack_loop["rounds"]) == 2
    assert attack_loop["rounds"][1]["round_kind"] == "retry"
    assert attack_loop["rounds"][1]["mode"] == "executed"
    assert attack_loop["rounds"][1]["expert_gate_decision"]["route_target"] == "delivery"
    assert attack_loop["loop_status"] == "baseline_executed_retry_executed_delivery_observation"
    assert data["delivery"]["sandbox_dispatcher"]["retry_attack"]["status"] == "executed"
    patch_result = next(
        item
        for item in data["delivery"]["backend_architecture"]["control_plane"]["results"]
        if item["stage"] == "patch_reflection"
    )
    assert (
        data["delivery"]["backend_architecture"]["control_plane"]["summary"]["same_run_retry"][
            "final_resolution"
        ]
        == "observed_after_retry"
    )
    assert patch_result["metadata"]["same_run_retry"]["final_resolution"] == "observed_after_retry"
    assert patch_result["decision_signal"] == "delivery"
    assert patch_result["decision_source"] == "retry_expert_gate_decision.route_target"
    assert patch_result["primary_output_ref"] == "delivery.attack_loop.retry_expert_gate_decision"

    timeline_response = client.get(f"/api/v1/cases/{case_id}/timeline")
    assert timeline_response.status_code == 200
    timeline = timeline_response.json()
    assert timeline["case_id"] == case_id
    assert timeline["latest_snapshot"]["metadata"]["retry_window_count"] == 3
    assert timeline["latest_snapshot"]["metadata"]["retry_handoff_count"] >= 3
    assert "attack_planning_agent:attack-plan-r2" in timeline["latest_snapshot"]["metadata"]["retry_projection_refs"]
    assert "expert_gate_agent:expert-gate-r2" in timeline["latest_snapshot"]["metadata"]["retry_projection_refs"]
    assert timeline["latest_snapshot"]["metadata"]["retry_compression_policies"] == [
        "retain_summary_and_refs_drop_raw_details"
    ]
    assert timeline["latest_snapshot"]["metadata"]["retry_resume_checkpoint_refs"]
    assert timeline["latest_snapshot"]["metadata"]["retry_resume_input_refs"]
    retry_checkpoint_ref = timeline["latest_snapshot"]["metadata"]["retry_resume_checkpoint_refs"][0]
    retry_input_ref = timeline["latest_snapshot"]["metadata"]["retry_resume_input_refs"][0]

    retry_checkpoint_snapshots_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/snapshots",
        params={"retry_resume_checkpoint_ref": retry_checkpoint_ref, "limit": 10},
    )
    assert retry_checkpoint_snapshots_response.status_code == 200
    retry_checkpoint_snapshots = retry_checkpoint_snapshots_response.json()
    assert retry_checkpoint_snapshots["case_id"] == case_id
    assert retry_checkpoint_snapshots["scope"]["retry_resume_checkpoint_ref"] == retry_checkpoint_ref
    assert retry_checkpoint_snapshots["total"] >= 1
    assert retry_checkpoint_snapshots["items"]

    retry_checkpoint_events_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/events",
        params={"retry_resume_checkpoint_ref": retry_checkpoint_ref, "limit": 10},
    )
    assert retry_checkpoint_events_response.status_code == 200
    retry_checkpoint_events = retry_checkpoint_events_response.json()
    assert retry_checkpoint_events["case_id"] == case_id
    assert retry_checkpoint_events["scope"]["retry_resume_checkpoint_ref"] == retry_checkpoint_ref
    assert retry_checkpoint_events["total"] >= 1
    assert retry_checkpoint_events["items"]

    retry_input_snapshots_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/snapshots",
        params={"retry_resume_input_ref": retry_input_ref, "limit": 10},
    )
    assert retry_input_snapshots_response.status_code == 200
    retry_input_snapshots = retry_input_snapshots_response.json()
    assert retry_input_snapshots["case_id"] == case_id
    assert retry_input_snapshots["scope"]["retry_resume_input_ref"] == retry_input_ref
    assert retry_input_snapshots["total"] >= 1
    assert retry_input_snapshots["items"]

    retry_input_events_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/events",
        params={"retry_resume_input_ref": retry_input_ref, "limit": 10},
    )
    assert retry_input_events_response.status_code == 200
    retry_input_events = retry_input_events_response.json()
    assert retry_input_events["case_id"] == case_id
    assert retry_input_events["scope"]["retry_resume_input_ref"] == retry_input_ref
    assert retry_input_events["total"] >= 1
    assert retry_input_events["items"]

    drilldown_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/drilldown",
        params={"projection_ref": "attack_planning_agent:attack-plan-r2", "limit": 10},
    )
    assert drilldown_response.status_code == 200
    drilldown = drilldown_response.json()
    assert drilldown["case_id"] == case_id
    assert drilldown["summary"]["retry_window_count"] == 3
    assert drilldown["summary"]["retry_handoff_count"] >= 3
    assert "attack_planning_agent:attack-plan-r2" in drilldown["summary"]["retry_projection_refs"]
    assert "vulnerability_agent:vulnerability-r2" in drilldown["summary"]["retry_projection_refs"]
    assert "expert_gate_agent:expert-gate-r2" in drilldown["summary"]["retry_projection_refs"]
    assert drilldown["summary"]["retry_typed_contract_refs"]
    assert drilldown["summary"]["retry_compression_policies"] == [
        "retain_summary_and_refs_drop_raw_details"
    ]
    assert drilldown["summary"]["retry_resume_checkpoint_refs"]
    assert drilldown["summary"]["retry_resume_input_refs"]

    retry_checkpoint_drilldown_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/drilldown",
        params={"retry_resume_checkpoint_ref": retry_checkpoint_ref, "limit": 10},
    )
    assert retry_checkpoint_drilldown_response.status_code == 200
    retry_checkpoint_drilldown = retry_checkpoint_drilldown_response.json()
    assert retry_checkpoint_drilldown["case_id"] == case_id
    assert retry_checkpoint_drilldown["scope"]["retry_resume_checkpoint_ref"] == retry_checkpoint_ref
    assert retry_checkpoint_drilldown["summary"]["matched_run_ids"]
    assert retry_checkpoint_drilldown["snapshots"]

    retry_input_drilldown_response = client.get(
        f"/api/v1/cases/{case_id}/timeline/drilldown",
        params={"retry_resume_input_ref": retry_input_ref, "limit": 10},
    )
    assert retry_input_drilldown_response.status_code == 200
    retry_input_drilldown = retry_input_drilldown_response.json()
    assert retry_input_drilldown["case_id"] == case_id
    assert retry_input_drilldown["scope"]["retry_resume_input_ref"] == retry_input_ref
    assert retry_input_drilldown["summary"]["matched_run_ids"]
    assert retry_input_drilldown["snapshots"]


def test_mas_execute_endpoint_can_disable_same_run_retry(monkeypatch):
    def fake_expert_gate(self, projection, *, run_id):  # noqa: ANN001
        target_service_ref = ""
        for artifact in projection.artifact_refs:
            metadata = artifact.metadata or {}
            target_service_ref = str(metadata.get("service_id") or "").strip()
            if target_service_ref:
                break
        return ExpertGateDecisionPayload(
            decision_id=f"expert-gate-{run_id[:8]}",
            target_service_ref=target_service_ref or "svc-retry-disabled",
            decision_family="retry_flow",
            decision_family_label="补充验证",
            action="retry_attack",
            action_label="补充攻击验证",
            route_target="attack_planning_agent",
            route_target_label="攻击重规划窗口",
            rationale="当前证据不足，理论上应继续补充攻击验证。",
            confidence=0.8,
            residual_risk_summary="当前残余风险仍需要额外采样确认。",
            follow_up_actions=["扩大攻击采样", "重新确认错误路径"],
        )

    monkeypatch.setattr(MASRuntimeSupport, "_decide_expert_gate_from_projection", fake_expert_gate)

    payload = {
        "requirement": "为支付系统设计兼顾性能与后量子安全的加密方案，延迟目标 20ms。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
        "max_same_run_retries": 0,
    }
    response = client.post("/api/v1/mas/execute", json=payload)
    assert response.status_code == 200

    data = response.json()
    attack_loop = data["delivery"]["attack_loop"]
    assert attack_loop["expert_gate_decision"]["route_target"] == "attack_planning_agent"
    assert attack_loop["retry_attack_decision"] == {}
    assert attack_loop["retry_expert_gate_decision"] == {}
    assert attack_loop["same_run_retry_budget"] == 0
    assert attack_loop["same_run_retry_used"] == 0
    assert attack_loop["same_run_retry_remaining"] == 0
    assert attack_loop["same_run_retry_summary"]["blocked_by_budget"] is True
    assert attack_loop["same_run_retry_summary"]["final_resolution"] == "blocked_by_budget"
    assert attack_loop["same_run_retry_summary"]["attempt_count"] == 1
    assert attack_loop["same_run_retry_summary"]["attempt_trace"][0]["status"] == "blocked"
    assert data["delivery"]["context_projections"]["attack_planning_retry"] == {}
    assert data["delivery"]["context_projections"]["vulnerability_evaluation_retry"] == {}
    memory_bus = data["delivery"]["backend_architecture"]["memory_bus"]
    assert memory_bus["summary"]["retry_window_count"] == 0
    assert memory_bus["summary"]["retry_handoff_count"] == 0
    assert len(attack_loop["rounds"]) == 1
    assert attack_loop["loop_status"] == "baseline_executed_retry_blocked_by_budget"
    assert (
        data["delivery"]["backend_architecture"]["control_plane"]["summary"]["same_run_retry"][
            "final_resolution"
        ]
        == "blocked_by_budget"
    )


def test_mas_execute_endpoint_retry_regate_can_enter_patch_flow(monkeypatch):
    original_plan_attack = MASRuntimeSupport._plan_attack_from_projection
    gate_calls = {"count": 0}

    def fake_expert_gate(self, projection, *, run_id):  # noqa: ANN001
        gate_calls["count"] += 1
        target_service_ref = ""
        for artifact in projection.artifact_refs:
            metadata = artifact.metadata or {}
            target_service_ref = str(metadata.get("service_id") or "").strip()
            if target_service_ref:
                break
        if gate_calls["count"] == 1:
            return ExpertGateDecisionPayload(
                decision_id=f"expert-gate-{run_id[:8]}",
                target_service_ref=target_service_ref or "svc-retry",
                decision_family="retry_flow",
                decision_family_label="补充验证",
                action="retry_attack",
                action_label="补充攻击验证",
                route_target="attack_planning_agent",
                route_target_label="攻击重规划窗口",
                rationale="当前基线证据不足，需要同轮补充攻击验证。",
                confidence=0.8,
                residual_risk_summary="当前残余风险仍需要额外采样确认。",
                follow_up_actions=["扩大攻击采样", "重新确认错误路径"],
            )
        return ExpertGateDecisionPayload(
            decision_id=f"expert-gate-{run_id[:8]}",
            target_service_ref=target_service_ref or "svc-retry",
            decision_family="patch_flow",
            decision_family_label="进入修补",
            action="plan_patch",
            action_label="进入修补",
            route_target="patch_agent",
            route_target_label="修补窗口",
            rationale="补充攻击证据已经足够支持进入 patch flow。",
            confidence=0.86,
            residual_risk_summary="需要尽快修补当前漏洞并做正式回归验证。",
            follow_up_actions=["生成修补方案", "执行正式回归"],
        )

    def fake_plan_attack(self, projection, *, run_id):  # noqa: ANN001
        if str(run_id).endswith("-retry"):
            _, baseline_specs = original_plan_attack(self, projection, run_id=run_id)
            return (
                AttackDecisionPayload(
                    decision_id=f"attack-decision-{run_id[:8]}",
                    target_service_ref="svc-retry",
                    action="execute",
                    action_label="执行补充攻击",
                    rationale="专家闸门要求同轮补充验证。",
                    selected_attack_family="oracle_probe",
                    expected_outcome="形成补充漏洞证据。",
                    confidence=0.79,
                    stop_conditions=["补充证据达到阈值"],
                    next_step="dispatch_attack",
                ),
                [
                    item.model_copy(update={"attack_id": f"{item.attack_id}-retry"})
                    for item in baseline_specs
                ],
            )
        return original_plan_attack(self, projection, run_id=run_id)

    monkeypatch.setattr(MASRuntimeSupport, "_decide_expert_gate_from_projection", fake_expert_gate)
    monkeypatch.setattr(MASRuntimeSupport, "_plan_attack_from_projection", fake_plan_attack)

    payload = {
        "requirement": "为支付系统设计兼顾性能与后量子安全的加密方案，延迟目标 20ms。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    response = client.post("/api/v1/mas/execute", json=payload)
    assert response.status_code == 200

    data = response.json()
    attack_loop = data["delivery"]["attack_loop"]
    assert attack_loop["expert_gate_decision"]["route_target"] == "attack_planning_agent"
    assert attack_loop["retry_expert_gate_decision"]["route_target"] == "patch_agent"
    assert attack_loop["patch_spec"]["patch_id"]
    assert attack_loop["patch_execution"]["execution_id"]
    assert attack_loop["same_run_retry_summary"]["final_resolution"] == "patched_after_retry"
    assert attack_loop["same_run_retry_summary"]["attempt_count"] == 1
    assert attack_loop["same_run_retry_summary"]["attempt_trace"][0]["status"] == "executed"
    assert attack_loop["same_run_retry_summary"]["attempt_trace"][0]["termination_reason"] == "patched_after_retry"
    assert data["delivery"]["context_projections"]["attack_planning_retry"]["agent_id"] == "attack_planning_agent"
    assert (
        data["delivery"]["context_projections"]["vulnerability_evaluation_retry"]["agent_id"]
        == "vulnerability_agent"
    )
    assert any(
        item["card_type"] == "retry_context_summary"
        for item in data["delivery"]["context_projections"]["attack_planning_retry"]["cards"]
    )
    memory_bus = data["delivery"]["backend_architecture"]["memory_bus"]
    assert memory_bus["summary"]["retry_window_count"] == 3
    assert "expert_gate_agent:expert-gate-r2" in memory_bus["summary"]["retry_projection_refs"]
    assert memory_bus["summary"]["retry_typed_contract_refs"]
    assert memory_bus["summary"]["retry_compression_policies"] == [
        "retain_summary_and_refs_drop_raw_details"
    ]
    assert memory_bus["summary"]["retry_resume_checkpoint_refs"]
    assert len(attack_loop["rounds"]) == 3
    assert attack_loop["rounds"][1]["round_kind"] == "retry"
    assert attack_loop["rounds"][2]["round_kind"] == "regression"
    assert attack_loop["loop_status"] == "baseline_executed_retry_executed_regression_executed"
    patch_result = next(
        item
        for item in data["delivery"]["backend_architecture"]["control_plane"]["results"]
        if item["stage"] == "patch_reflection"
    )
    assert patch_result["decision_signal"] == "replan"
    assert patch_result["decision_source"] == "regression_attack_decision.action_or_patch_execution.status"
    assert patch_result["primary_output_ref"] == "delivery.attack_loop.patch_spec"


def test_case_memory_endpoint_returns_persisted_snapshot():
    case_id = f"case-read-{uuid4().hex[:8]}"
    payload = {
        "requirement": "面向建筑项目 CDE 设计具备后量子迁移能力的 BIM/IFC 审计型签名与加密方案。",
        "case_id": case_id,
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    execute_response = client.post("/api/v1/mas/execute?use_langgraph=true", json=payload)
    assert execute_response.status_code == 200

    read_response = client.get(f"/api/v1/cases/{case_id}")
    assert read_response.status_code == 200
    data = read_response.json()
    assert data["case_id"] == case_id
    assert "decision_log" in data
    assert "recent_reflections" in data
    assert data["recent_reflections"]
    assert "status_label" in data


def test_case_memory_list_endpoint_returns_recent_summaries():
    case_id = f"case-list-{uuid4().hex[:8]}"
    payload = {
        "requirement": "面向建筑工程协同平台设计支持 ISO 19650 证据审计与后量子迁移的加密方案。",
        "case_id": case_id,
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    execute_response = client.post("/api/v1/mas/execute?use_langgraph=true", json=payload)
    assert execute_response.status_code == 200

    list_response = client.get("/api/v1/cases?limit=10")
    assert list_response.status_code == 200
    data = list_response.json()
    assert data["total"] >= 1

    matched = next((item for item in data["items"] if item["case_id"] == case_id), None)
    assert matched is not None
    assert "status_label" in matched
    assert matched["reflection_count"] >= 1
    assert "blocking_count" in matched
    assert "decision_count" in matched


def test_case_memory_delete_endpoint_removes_snapshot():
    case_id = f"case-delete-{uuid4().hex[:8]}"
    payload = {
        "requirement": "面向建筑工程协作平台设计支持审计追踪的后量子迁移方案。",
        "case_id": case_id,
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    execute_response = client.post("/api/v1/mas/execute?use_langgraph=true", json=payload)
    assert execute_response.status_code == 200

    delete_response = client.delete(f"/api/v1/cases/{case_id}")
    assert delete_response.status_code == 200
    deleted = delete_response.json()
    assert deleted["case_id"] == case_id
    assert deleted["deleted"] is True

    read_response = client.get(f"/api/v1/cases/{case_id}")
    assert read_response.status_code == 404


def test_knowledge_ingestion_endpoint_accepts_docx_upload(tmp_path):
    source_path = tmp_path / "sample_policy.docx"
    doc = Document()
    doc.add_heading("3.2.1 密钥轮换要求", level=1)
    doc.add_paragraph("敏感数据平台应明确密钥轮换周期，并保留审计留痕。")
    doc.save(source_path)

    with source_path.open("rb") as handle:
        response = client.post(
            "/api/v1/knowledge/ingest",
            data={
                "doc_type": "policy",
                "metadata_json": "{\"region\":\"CN\",\"industry\":\"construction\"}",
                "upsert_qdrant": "false",
            },
            files={
                "files": (
                    source_path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert response.status_code == 200
    data = response.json()
    assert data["doc_type"] == "policy"
    assert data["total_files"] == 1
    assert data["total_chunks"] >= 1
    output_path = Path(data["output_path"])
    assert output_path.exists()
    assert data["files"][0]["file_name"] == source_path.name
    assert data["files"][0]["chunk_count"] >= 1
    assert data["files"][0]["previews"]

    stored_path = Path(data["files"][0]["stored_path"])
    if output_path.exists():
        output_path.unlink()
    if stored_path.exists():
        stored_path.unlink()
        if stored_path.parent.exists():
            stored_path.parent.rmdir()


def test_knowledge_ingestion_delete_endpoint_removes_local_artifacts(tmp_path):
    source_path = tmp_path / "sample_case.docx"
    doc = Document()
    doc.add_heading("整改建议", level=1)
    doc.add_paragraph("这里包含演示用的企业整改记录，删除接口应清理本地上传文件与 JSONL。")
    doc.save(source_path)

    with source_path.open("rb") as handle:
        ingest_response = client.post(
            "/api/v1/knowledge/ingest",
            data={
                "doc_type": "case",
                "metadata_json": "{\"region\":\"CN\",\"industry\":\"construction\"}",
                "upsert_qdrant": "false",
            },
            files={
                "files": (
                    source_path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert ingest_response.status_code == 200
    ingested = ingest_response.json()
    request_id = ingested["request_id"]
    output_path = Path(ingested["output_path"])
    stored_path = Path(ingested["files"][0]["stored_path"])

    assert output_path.exists()
    assert stored_path.exists()

    delete_response = client.delete(f"/api/v1/knowledge/ingest/{request_id}")
    assert delete_response.status_code == 200
    deleted = delete_response.json()
    assert deleted["request_id"] == request_id
    assert deleted["deleted"] is True
    assert deleted["deleted_paths"]

    assert not output_path.exists()
    assert not stored_path.exists()
    assert not stored_path.parent.exists()

    list_response = client.get("/api/v1/knowledge/ingestions?limit=10")
    assert list_response.status_code == 200
    listed = list_response.json()
    matched = next((item for item in listed["items"] if item["request_id"] == request_id), None)
    assert matched is not None
    assert matched["local_artifacts_present"] is False
    assert matched["deleted_at"] is not None


def test_knowledge_ingestion_list_endpoint_returns_recent_records(tmp_path):
    source_path = tmp_path / "sample_standard.docx"
    doc = Document()
    doc.add_heading("7.1 审计留痕要求", level=1)
    doc.add_paragraph("关键数据访问和密钥操作应保留可追溯审计记录。")
    doc.save(source_path)

    with source_path.open("rb") as handle:
        ingest_response = client.post(
            "/api/v1/knowledge/ingest",
            data={
                "doc_type": "standard",
                "metadata_json": "{\"region\":\"CN\",\"industry\":\"construction\",\"scenario\":\"bim_delivery\"}",
                "upsert_qdrant": "false",
            },
            files={
                "files": (
                    source_path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert ingest_response.status_code == 200
    ingested = ingest_response.json()
    request_id = ingested["request_id"]

    list_response = client.get("/api/v1/knowledge/ingestions?limit=10")
    assert list_response.status_code == 200
    listed = list_response.json()
    assert listed["total"] >= 1

    matched = next((item for item in listed["items"] if item["request_id"] == request_id), None)
    assert matched is not None
    assert matched["doc_type"] == "standard"
    assert matched["total_files"] == 1
    assert matched["total_chunks"] >= 1
    assert matched["local_artifacts_present"] is True
    assert matched["metadata"]["region"] == "CN"
    assert matched["files"]

    delete_response = client.delete(f"/api/v1/knowledge/ingest/{request_id}")
    assert delete_response.status_code == 200


def test_knowledge_ingestion_qdrant_delete_endpoint_updates_manifest(tmp_path, monkeypatch):
    source_path = tmp_path / "sample_qdrant.docx"
    doc = Document()
    doc.add_heading("5.4 密钥托管要求", level=1)
    doc.add_paragraph("该文档用于测试按 request_id 清理 Qdrant 副本后的 manifest 回写。")
    doc.save(source_path)

    with source_path.open("rb") as handle:
        ingest_response = client.post(
            "/api/v1/knowledge/ingest",
            data={
                "doc_type": "policy",
                "metadata_json": "{\"region\":\"CN\",\"industry\":\"construction\"}",
                "upsert_qdrant": "false",
            },
            files={
                "files": (
                    source_path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert ingest_response.status_code == 200
    ingested = ingest_response.json()
    request_id = ingested["request_id"]

    service = KnowledgeIngestionService()
    service._update_record(
        request_id,
        qdrant_requested=True,
        qdrant_upserted=True,
        qdrant_cleanup_required=True,
        qdrant_message="已写入 Qdrant：1 条知识块",
    )

    monkeypatch.setattr(
        KnowledgeIngestionService,
        "_delete_qdrant_points",
        lambda self, record: record.total_chunks,
    )

    delete_response = client.delete(f"/api/v1/knowledge/ingest/{request_id}/qdrant")
    assert delete_response.status_code == 200
    deleted = delete_response.json()
    assert deleted["request_id"] == request_id
    assert deleted["deleted"] is True
    assert deleted["deleted_points"] >= 1
    assert deleted["qdrant_deleted_at"] is not None

    list_response = client.get("/api/v1/knowledge/ingestions?limit=10")
    assert list_response.status_code == 200
    listed = list_response.json()
    matched = next((item for item in listed["items"] if item["request_id"] == request_id), None)
    assert matched is not None
    assert matched["qdrant_upserted"] is False
    assert matched["qdrant_cleanup_required"] is False
    assert matched["qdrant_deleted_at"] is not None

    cleanup_response = client.delete(f"/api/v1/knowledge/ingest/{request_id}")
    assert cleanup_response.status_code == 200


def test_knowledge_ingestion_qdrant_reingest_endpoint_updates_manifest(tmp_path, monkeypatch):
    source_path = tmp_path / "sample_reingest.docx"
    doc = Document()
    doc.add_heading("6.2 审计日志要求", level=1)
    doc.add_paragraph("该文档用于测试按 request_id 从本地 JSONL 重建 Qdrant 副本。")
    doc.save(source_path)

    with source_path.open("rb") as handle:
        ingest_response = client.post(
            "/api/v1/knowledge/ingest",
            data={
                "doc_type": "policy",
                "metadata_json": "{\"region\":\"CN\",\"industry\":\"construction\"}",
                "upsert_qdrant": "false",
            },
            files={
                "files": (
                    source_path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert ingest_response.status_code == 200
    ingested = ingest_response.json()
    request_id = ingested["request_id"]

    service = KnowledgeIngestionService()
    service._update_record(
        request_id,
        qdrant_requested=False,
        qdrant_upserted=False,
        qdrant_cleanup_required=False,
        qdrant_deleted_at=datetime(2026, 3, 25, tzinfo=timezone.utc),
        qdrant_message="Qdrant 副本已删除",
    )

    monkeypatch.setattr(
        KnowledgeIngestionService,
        "_upsert_qdrant_payloads",
        lambda self, payloads: len(payloads),
    )

    reingest_response = client.post(f"/api/v1/knowledge/ingest/{request_id}/qdrant")
    assert reingest_response.status_code == 200
    rebuilt = reingest_response.json()
    assert rebuilt["request_id"] == request_id
    assert rebuilt["reingested"] is True
    assert rebuilt["upserted_points"] >= 1

    list_response = client.get("/api/v1/knowledge/ingestions?limit=10")
    assert list_response.status_code == 200
    listed = list_response.json()
    matched = next((item for item in listed["items"] if item["request_id"] == request_id), None)
    assert matched is not None
    assert matched["qdrant_requested"] is True
    assert matched["qdrant_upserted"] is True
    assert matched["qdrant_cleanup_required"] is True
    assert matched["qdrant_deleted_at"] is None

    cleanup_response = client.delete(f"/api/v1/knowledge/ingest/{request_id}")
    assert cleanup_response.status_code == 200


def test_mas_stream_endpoint():
    payload = {
        "requirement": "为物联网设计低延迟加密方案，并兼顾未来量子安全。",
        "num_variants": 2,
        "generate_code": False,
        "max_audit_rounds": 2,
    }
    with client.stream("POST", "/api/v1/mas/stream", json=payload) as response:
        assert response.status_code == 200
        lines = [line for line in response.iter_lines() if line]

    assert any('"type": "progress"' in line for line in lines)
    assert any('"type": "final"' in line for line in lines)
    assert any('"engine": "langgraph"' in line for line in lines if '"type": "run"' in line)


def test_cancel_unknown_run():
    response = client.post("/api/v1/mas/cancel/non-existent-run")
    assert response.status_code == 404


def test_mas_report_endpoint():
    execute_payload = {
        "requirement": "建筑工程 BIM/IFC 交付需要内容防篡改、批准版本校验和长期后量子迁移路线。",
        "num_variants": 3,
        "generate_code": True,
        "max_audit_rounds": 3,
    }
    execute_resp = client.post("/api/v1/mas/execute", json=execute_payload)
    assert execute_resp.status_code == 200
    mas_result = execute_resp.json()
    discussion_log = mas_result["discussion_log"]
    assert discussion_log
    assert all(item["actor"] in STABLE_ACTORS for item in discussion_log)
    assert any(contains_chinese(item.get("actor_label", "")) for item in discussion_log)
    assert any(contains_chinese(item["message"]) for item in discussion_log)
    assert any(contains_chinese(item.get("status_label", "")) for item in discussion_log)

    credibility = mas_result["credibility_assessment"]
    assert credibility["trust_level"] in {"high", "medium", "low", "unknown"}
    assert contains_chinese(credibility.get("trust_level_label", ""))
    assert credibility.get("component_evidence")
    assert any(
        contains_chinese(item.get("support_level_label", ""))
        for item in credibility.get("component_evidence", [])
        if isinstance(item, dict)
    )
    strengths_and_gaps = credibility.get("strengths", []) + credibility.get("gaps", [])
    assert strengths_and_gaps
    assert any(contains_chinese(item) for item in strengths_and_gaps)

    production_guide = mas_result["delivery"]["production_guide"]
    assert production_guide
    assert all(contains_chinese(item) for item in production_guide)
    assert contains_chinese(mas_result["delivery"].get("status_label", ""))

    auditor_rounds = mas_result["auditor_rounds"]
    assert auditor_rounds
    assert any(contains_chinese(item.get("verdict_label", "")) for item in auditor_rounds)
    assert all(item.get("quantum_ready") in {True, False} for item in auditor_rounds)
    assert all(contains_chinese(item.get("quantum_ready_label", "")) for item in auditor_rounds)

    attempts = mas_result["engineer"]["attempts"]
    assert attempts
    assert any(contains_chinese(item.get("status_label", "")) for item in attempts)

    variant_comparison = mas_result["delivery"]["variant_comparison"]
    if variant_comparison.get("available"):
        recommendations = variant_comparison["comparison"].get("recommendations", [])
        assert recommendations
        assert any(not contains_chinese(item.get("use_case", "")) for item in recommendations)
        assert any(contains_chinese(item.get("use_case_label", "")) for item in recommendations)
        assert any(contains_chinese(item.get("reason", "")) for item in recommendations)
        side_by_side = variant_comparison["comparison"].get("side_by_side_table", [])
        assert side_by_side
        assert all(isinstance(row.get("quantum_resistant"), bool) for row in side_by_side)
        assert all(row.get("quantum_resistant_label") in {"是", "否", "--"} for row in side_by_side)
    else:
        assert contains_chinese(variant_comparison.get("note", ""))

    compliance_report = mas_result["compliance_report"]
    assert contains_chinese(compliance_report.get("summary_text", ""))
    assert any(
        contains_chinese(item.get("recommendation", ""))
        for item in compliance_report.get("recommendations", [])
        if isinstance(item, dict)
    )

    vulnerability_report = mas_result["vulnerability_report"]
    assert contains_chinese(vulnerability_report.get("summary_text", ""))
    if vulnerability_report.get("recommendations"):
        assert any(
            contains_chinese(item.get("action", ""))
            for item in vulnerability_report.get("recommendations", [])
            if isinstance(item, dict)
        )

    report_payload = {
        "mas_result": mas_result,
        "scenario": "construction",
        "include_code": True,
        "include_pdf": False,
    }
    report_resp = client.post("/api/v1/mas/report", json=report_payload)
    assert report_resp.status_code == 200
    data = report_resp.json()
    assert data["scenario"] == "construction"
    assert data["selected_scheme"]
    assert data["template_id"] == "construction_trusted_delivery"
    assert data["template_name"]
    assert data["enterprise_delivery"]["sections"]
    assert "<!doctype html>" in data["html"].lower()
    assert "BuildTrust 企业交付报告" in data["html"]
    assert "可信度概览" in data["html"]
    assert "企业交付模板" in data["html"]
    assert "执行摘要" in data["markdown"]
    assert "可信度概览" in data["markdown"]
    assert "企业交付模板" in data["markdown"]
    assert "\\section{执行摘要}" in data["latex"]
    assert "\\section{可信度概览}" in data["latex"]
    assert "\\section{企业交付模板}" in data["latex"]
    assert "合规与风险摘要" in data["html"]
    assert "合规与风险摘要" in data["markdown"]
    assert "\\section{合规与风险摘要}" in data["latex"]
    assert "合规建议：" in data["html"] or "漏洞处置建议：" in data["html"]
    assert "合规建议：" in data["markdown"] or "漏洞处置建议：" in data["markdown"]
    assert isinstance(data["deployment_guide"], list)

def test_env_settings_roundtrip():
    response = client.get("/api/v1/settings/env")
    assert response.status_code == 200
    assert "values" in response.json()


def test_llm_validate_unknown_provider():
    response = client.post(
        "/api/v1/llm/validate",
        json={"provider": "unknown-provider", "prompt": "OK"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["success"] is False


def test_structured_spec_detects_chinese_construction_quantum_keywords():
    service = MASOrchestrationService()
    requirement = build_requirement(quantum_resistant=False, security_level=128)
    raw_text = "\u5efa\u7b51\u9879\u76ee BIM/IFC \u4ea4\u4ed8\u9700\u8981\u4f4e\u5ef6\u8fdf\u7b7e\u540d\u4e0e\u52a0\u5bc6\u5e76\u6ee1\u8db3 ISO 19650\uff0c\u8981\u6c42\u540e\u91cf\u5b50\u53ef\u8fc1\u79fb\u3002"

    structured_spec = service._build_structured_spec(requirement, raw_text)

    assert structured_spec["domain"] == "construction"
    assert structured_spec["compliance"] == "ISO_19650"
    assert structured_spec["quantum_safe"] is True


def test_audit_input_normalizes_hybrid_pqc_scheme_for_compliance():
    service = MASOrchestrationService()
    requirement = build_requirement(quantum_resistant=True, security_level=256)
    scheme = next(
        item
        for item in service._heuristic_generate_schemes(requirement, 3)
        if {"AES", "GCM"}.issubset({component.name for component in item.architecture.components})
    )

    audit_input = service._scheme_to_audit_input(scheme)
    standards = service._select_standards(
        {"domain": "construction", "compliance": "ISO_19650", "quantum_safe": True}
    )
    normalized_score = service.compliance_reporter.generate_report(audit_input, standards)["overall_compliance"]
    legacy_score = service.compliance_reporter.generate_report(
        {**audit_input, "algorithm": "x25519", "mode": ""},
        standards,
    )["overall_compliance"]

    assert audit_input["algorithm"] == "aes-gcm"
    assert audit_input["mode"] == "gcm"
    assert "kyber" in audit_input["quantum_assessment_input"]
    assert standards == [
        ComplianceStandard.NIST_CSF,
        ComplianceStandard.ISO_27001,
        ComplianceStandard.FIPS_140_3,
    ]
    assert normalized_score > legacy_score


def test_prepare_schemes_for_audit_hardens_quantum_requirement_and_reranks():
    service = MASOrchestrationService()
    requirement = build_requirement(quantum_resistant=True, security_level=256)
    structured_spec = {"domain": "construction", "compliance": "ISO_19650", "quantum_safe": True}
    schemes = [
        build_scheme(service, "proposal-chacha", ["ChaCha20-Poly1305", "AES"], score=9.8),
        build_scheme(service, "proposal-mixed", ["AES", "ChaCha20-Poly1305"], score=9.2),
        build_scheme(service, "proposal-aes-gcm", ["AES", "GCM"], score=8.1),
    ]

    prepared = service._prepare_schemes_for_audit(schemes, requirement, structured_spec)
    first_components = {item.name for item in prepared[0].architecture.components}
    first_audit_input = service._scheme_to_audit_input(prepared[0])

    assert all(
        any("Kyber" in component.name for component in scheme.architecture.components)
        for scheme in prepared
    )
    assert prepared[0].metadata.name == "proposal-aes-gcm"
    assert "Kyber-768" in first_components or "Kyber-1024" in first_components
    assert first_audit_input["algorithm"] == "aes-gcm"
